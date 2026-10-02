"""Full 72-task DocuBench adapter, including its seven native office/text tasks."""
import argparse
import hashlib
import importlib.util
import json
import sys
import time
import xml.etree.ElementTree as ET
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from docquorum import config as C
from docquorum.extract import extract
from docquorum.tools import llm
from docquorum.agents import prompts
from docquorum.checks.conform import conform
from resource_profile import install
from varex import write

SOURCE = ROOT / "data/multibench_sources/docubench"
OUT = ROOT / "results/multibench/docubench"


def native_text(path):
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        formulas = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
        lines = []
        for sheet in formulas:
            lines.append(f"WORKSHEET {sheet.title}")
            for row in sheet:
                cells = []
                for cell in row:
                    if cell.value is not None:
                        value = cached[sheet.title][cell.coordinate].value
                        cells.append(f"{cell.coordinate}: {value if value is not None else cell.value}")
                if cells: lines.append(" | ".join(cells))
        formulas.close(); cached.close()
        return "\n".join(lines)
    if path.suffix.lower() == ".docx":
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(path) as archive:
            names = ["word/document.xml"] + sorted(n for n in archive.namelist()
                if n.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes")) and n.endswith(".xml"))
            lines = []
            for name in names:
                lines.append(name)
                tree = ET.fromstring(archive.read(name))
                for para in tree.iter(ns + "p"):
                    lines.append("".join(element.text or "" for element in para.iter(ns + "t")))
            return "\n".join(lines)
    # Raw HTML/XML retain their tags, attributes, table cells and ordering.
    return path.read_text(encoding="utf-8-sig")


def records():
    docs = {p.stem: p for p in (SOURCE / "documents").iterdir() if p.is_file()}
    schemas = {p.stem: p for p in (SOURCE / "schemas").glob("*.json")}
    if set(docs) != set(schemas) or len(docs) != 72:
        raise ValueError("DocuBench must contain 72 complete input/schema pairs")
    rows = []
    for name in sorted(docs):
        guideline = SOURCE / "guidelines" / f"{name}.txt"
        rows.append({"id": name, "path": str(docs[name]), "format": docs[name].suffix.lower(),
                     "source_sha256": hashlib.sha256(docs[name].read_bytes()).hexdigest(),
                     "schema_sha256": hashlib.sha256(schemas[name].read_bytes()).hexdigest(),
                     "guidelines_sha256": hashlib.sha256(guideline.read_bytes()).hexdigest() if guideline.exists() else None})
    return rows


def run(workers, version="bounded-v1"):
    installer = install
    if version == 'bounded-v2':
        from resource_profile_v2 import install as installer
    profile = installer()
    manifest = {"source_commit": "43a3f3bc00e591e711075678ca6d154acfedcf42", "version": "0.2.0",
                "model": C.MODEL, "profile": profile, "documents": records()}
    existing = OUT / "manifest.json"
    if existing.exists() and json.loads(existing.read_text()) != manifest:
        raise ValueError("Input or run configuration changed; use a new run directory")
    write(existing, manifest)

    def one(row):
        name, doc = row["id"], Path(row["path"])
        dst = OUT / "predictions" / f"{name}.json"
        if dst.exists(): return name
        start = time.monotonic()
        schema = json.loads((SOURCE / "schemas" / f"{name}.json").read_text())
        guideline = SOURCE / "guidelines" / f"{name}.txt"
        task = guideline.read_text() if guideline.exists() else None
        try:
            if doc.suffix.lower() in {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
                if doc.suffix.lower() in {".tif", ".tiff"}:
                    from PIL import Image, ImageSequence
                    pages = [frame.convert("RGB") for frame in ImageSequence.Iterator(Image.open(doc))]
                    converted = OUT / "converted" / f"{name}.pdf"
                    converted.parent.mkdir(parents=True, exist_ok=True)
                    pages[0].save(converted, "PDF", save_all=True, append_images=pages[1:], resolution=150)
                    doc = converted
                result, report = extract(doc, schema, task=task, workers=2,
                    log=lambda msg: print(f"{name}: {msg}", flush=True))
            else:
                text = native_text(doc)
                prompt = prompts.WHOLE.format(rules="", task=f"TASK\n{task or ''}\n\n", n=1,
                    modality=f"Its complete native {doc.suffix} source is below.",
                    text_block=prompts.TEXT_LAYER.format(body=text), schema=json.dumps(schema, indent=1))
                cl = llm.client()
                response = cl.chat.completions.create(model=C.MODEL, temperature=0, max_tokens=8192,
                    response_format={"type": "json_object"}, extra_body=llm.extra_body(),
                    messages=[{"role": "user", "content": prompt}])
                choice = response.choices[0]
                if choice.finish_reason != "stop":
                    raise ValueError("Native-format response was truncated")
                result = json.loads(choice.message.content)
                if not isinstance(result, dict): raise ValueError("Expected a JSON object")
                result = conform(result, schema, schema)
                report = {"readings": ["native_text"], "native_format": row["format"], "requests": cl.calls}
            info = {"status": "completed", "report": report}
        except Exception as error:
            result, info = {}, {"status": "failed", "error": f"{type(error).__name__}: {str(error)[:500]}"}
        info["seconds"] = time.monotonic() - start
        write(OUT / "reports" / f"{name}.json", info)
        write(dst, {"data": result, "meta": {"engine": profile['profile'], "status": info["status"], "model": C.MODEL}})
        return name
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = [pool.submit(one, row) for row in manifest["documents"]]
        for i, f in enumerate(as_completed(futures), 1):
            print(f"[{i}/72] {f.result()}", flush=True)


def score():
    spec = importlib.util.spec_from_file_location("docubench_official_scorer", SOURCE / "scorer.py")
    scorer = importlib.util.module_from_spec(spec); spec.loader.exec_module(scorer)
    rows = []
    for row in records():
        name = row["id"]
        path = OUT / "predictions" / f"{name}.json"
        status = "missing"
        if path.exists():
            payload = json.loads(path.read_text())
            status = payload["meta"]["status"]
        if status == "completed":
            value = scorer.score_standardization(result=payload["data"],
                schema=json.loads((SOURCE / "schemas" / f"{name}.json").read_text()),
                label=json.loads((SOURCE / "labels" / f"{name}.json").read_text()))["final"]
        else:
            value = 0.0
        rows.append({"id": name, "format": row["format"], "status": status, "score": value})
    manifest = json.loads((OUT / 'manifest.json').read_text())
    report = {"benchmark": "DocuBench", "version": "0.2.0", "engine": manifest['profile']['profile'],
              "metric": "macro_average_field_accuracy", "score": 100*sum(r["score"] for r in rows)/72,
              "expected": 72, "attempted": sum(r["status"] != "missing" for r in rows),
              "completed": sum(r["status"] == "completed" for r in rows), "per_document": rows}
    report["eligible_for_full_ranking"] = report["attempted"] == 72
    write(OUT / "summary.json", report)
    print(json.dumps({k:v for k,v in report.items() if k != "per_document"}, indent=2))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("action", choices=["run", "score"])
    ap.add_argument("--workers", type=int, default=2)
    ap.add_argument("--out", type=Path, default=OUT)
    ap.add_argument("--profile", choices=['bounded-v1', 'bounded-v2'], default='bounded-v1')
    a = ap.parse_args()
    OUT = a.out.resolve()
    if a.action == "run": run(a.workers, a.profile)
    else: score()
